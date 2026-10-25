import random
import time

from openpyxl import load_workbook

from Tapi_Project_Solver_Engine import (
    _read_input_from_workbook,
    _write_results_workbook,
    solve_lp_subproblem,
)
#  הגדרת מגבלות האלגוריתם הגנטי (לפי הדרישות)
POPULATION_MIN = 20
POPULATION_MAX = 200
ELITE_MAX_FRACTION = 0.10


def read_ga_parameters(input_file_name):
    # קריאת פרמטרי ההרצה של האלגוריתם הגנטי מתאי הקלט באקסל
    file_path = f"{input_file_name}.xlsx"
    workbook = load_workbook(file_path, data_only=True)
    sheet = workbook.active

    parameters = {
        "population_size": sheet["A7"].value,
        "elite_size": sheet["A9"].value,
        "mutation_prob": sheet["A11"].value,
        "time_limit": sheet["A13"].value,
    }

    workbook.close()
    return parameters


def _parse_even_integer(value, field_name):
# פונקציית עזר לווידוא שהפרמטרים שהוזנו הם מספרים שלמים וזוגיים
    if value is None or value == "":
        return False, None, f"{field_name} is missing."

    if isinstance(value, bool):
        return False, None, f"{field_name} must be an integer, not a boolean."

    try:
        numeric_value = float(value)
    except (TypeError, ValueError):
        return False, None, f"{field_name} must be a numeric value (got {value!r})."

    if numeric_value != int(numeric_value):
        return False, None, f"{field_name} must be an integer (got {value!r})."

    parsed_value = int(numeric_value)
    if parsed_value % 2 != 0:
        return False, None, f"{field_name} must be an even integer (got {parsed_value})."

    return True, parsed_value, ""


def validate_ga_parameters(population_size, elite_size):
  # ביצוע בדיקות תקינות לפרמטרים
    errors = []

    pop_valid, pop_value, pop_error = _parse_even_integer(
        population_size, "Population size"
    )
    if not pop_valid:
        errors.append(pop_error)
    elif not (POPULATION_MIN <= pop_value <= POPULATION_MAX):
        errors.append(
            f"Population size must be between {POPULATION_MIN} and {POPULATION_MAX} "
            f"(got {pop_value})."
        )

    elite_valid, elite_value, elite_error = _parse_even_integer(
        elite_size, "Elite solutions count"
    )
    if not elite_valid:
        errors.append(elite_error)
    elif pop_valid:
        max_elite = int(pop_value * ELITE_MAX_FRACTION)
        if elite_value > max_elite:
            errors.append(
                "Elite solutions count must be at most 10% of the population size "
                f"(population={pop_value}, maximum allowed elite={max_elite}, got {elite_value})."
            )

    return len(errors) == 0, errors


def read_job_data(input_file_name):
# שליפת נתוני המשימות (זמני עיבוד, זמני יעד וקנסות) דרך פונקציית העזר של קובץ הסולבר
    file_path = f"{input_file_name}.xlsx"
    workbook = load_workbook(file_path, data_only=True)
    sheet = workbook.active
    t, d, v, w, job_labels = _read_input_from_workbook(sheet)
    workbook.close()
    return t, d, v, w, job_labels


def create_initial_population(n_jobs, population_size):
  #  הגרלת אוכלוסייה ראשונית של סידורי משימות
    job_indices = list(range(n_jobs))
    population = []

    for _ in range(population_size):
        individual = random.sample(job_indices, n_jobs)
        population.append(individual)

    return population


def evaluate_population(population, t, d, v, w):
# מדידת איכות הפתרונות לכל סידור משימות מופעל הסולבר לשיבוץ זמני בטלה בצורה מיטבית
    scores = []

    for sequence in population:
        obj_value, _, _, _ = solve_lp_subproblem(sequence, t, d, v, w)
        scores.append(obj_value)

    return scores


def _score_key(score):
  # פונקציית מיון לבעיית מזעור- פתרון שלא עבד מקבל ציון אינסוף ועובר לסוף הרשימה
    return float("inf") if score is None else score


def calculate_ranking_probabilities(scores):
#  קביעת הסתברויות בשיטת Ranking  
# n סידור בסדר יורד כך שהפתרון הגרוע מקבל דירוג 1, והטוב ביותר מקבל דירוג 
    n = len(scores)
    ranked_indices = sorted(range(n), key=lambda idx: _score_key(scores[idx]), reverse=True)
    ranks = [0] * n

    for rank, idx in enumerate(ranked_indices, start=1):
        ranks[idx] = rank

    total_rank = n * (n + 1) / 2
    return [rank / total_rank for rank in ranks]


def select_parent(population, probabilities):
# הגרלת הורה מתוך האוכלוסייה בהתאם להסתברויות שחושבו
    return random.choices(population, weights=probabilities, k=1)[0]


def order_crossover(parent1, parent2):
  #  העברת רצף גנטי בלי לייצר משימות כפולות
    n = len(parent1)
    start, end = sorted(random.sample(range(n), 2))
    child = [None] * n

    # שמירת מקטע רציף מההורה הראשון
    child[start : end + 1] = parent1[start : end + 1]

# השלמת שאר המשימות מההורה השני לפי סדר הופעתן
    fill_values = [gene for gene in parent2 if gene not in child]
    fill_positions = [i for i in range(n) if child[i] is None]

    for position, gene in zip(fill_positions, fill_values):
        child[position] = gene

    return child


def crossover(parent1, parent2):
# זיווג הורים ליצירת שני צאצאים 
    child1 = order_crossover(parent1, parent2)
    child2 = order_crossover(parent2, parent1)
    return child1, child2


def mutate(individual, mutation_prob):
# שימוש במוטציות בהתאם לכמות המשימות, כדי להקטין סיכוי להתכנסות מוקדמת
    if random.random() < mutation_prob:
        num_swaps = random.randint(1, max(2, len(individual) // 5))
        for _ in range(num_swaps):
            i, j = random.sample(range(len(individual)), 2)
           # ביצוע החלפה בפועל בין המשימות
            individual[i], individual[j] = individual[j], individual[i]
    return individual


def select_elite(population, scores, elite_size):
# העברת הפתרונות הטובים כפי שהם לדור הבא
    ranked = sorted(range(len(scores)), key=lambda idx: _score_key(scores[idx]))
    elite_indices = ranked[:elite_size]

    elite_population = [population[idx][:] for idx in elite_indices]
    elite_scores = [scores[idx] for idx in elite_indices]
    return elite_population, elite_scores


def build_next_generation(
    population, scores, t, d, v, w, population_size, elite_size, mutation_prob
):
# בניית הדור הבא 
    probabilities = calculate_ranking_probabilities(scores)
    elite_population, elite_scores = select_elite(population, scores, elite_size)

    offspring = []
    offspring_needed = population_size - elite_size
# השלמת אוכלוסייה לדור הבא באמצעות זיווג ומוטציות
    while len(offspring) < offspring_needed:
        parent1 = select_parent(population, probabilities)
        parent2 = select_parent(population, probabilities)
        child1, child2 = crossover(parent1, parent2)

        offspring.append(mutate(child1, mutation_prob))
        if len(offspring) < offspring_needed:
            offspring.append(mutate(child2, mutation_prob))
# חישוב ערכי פונקציית המטרה רק לצאצאים החדשים שנוספו
    offspring_scores = evaluate_population(offspring, t, d, v, w)
    next_population = elite_population + offspring
    next_scores = elite_scores + offspring_scores
    return next_population, next_scores


def _update_best_solution(best_record, sequence, score, generation):
# שמירת הפתרון הטוב ביותר שנמצא בכל הריצה
    if score is None:
        return best_record

    if best_record["score"] is None or score < best_record["score"]:
        best_record["sequence"] = sequence[:]
        best_record["score"] = score
        best_record["generation"] = generation

    return best_record


def _population_best(population, scores):
    # שליפת המשתמש הטוב ביותר באוכלוסייה הנוכחית
    best_idx = min(range(len(scores)), key=lambda idx: _score_key(scores[idx]))
    return population[best_idx][:], scores[best_idx]


def Run_GA(Input_File_Name, Output_File_Na):
 # פונקצייה המנהלת את כלל ריצת האלגוריתם הגנטי מהקלט ועד להפקת הגאנט
    try:
        ga_parameters = read_ga_parameters(Input_File_Name)
    except FileNotFoundError:
        print(f"Error: Input file '{Input_File_Name}.xlsx' was not found.")
        return

    is_valid, error_messages = validate_ga_parameters(
        ga_parameters["population_size"],
        ga_parameters["elite_size"],
    )

    if not is_valid:
        print("GA input validation failed:")
        for message in error_messages:
            print(f"  - {message}")
        return

    population_size = int(float(ga_parameters["population_size"]))
    elite_size = int(float(ga_parameters["elite_size"]))
    mutation_prob = float(ga_parameters["mutation_prob"])
    time_limit = float(ga_parameters["time_limit"])

    print("GA input validation passed.")
    print(f"  Population size      : {population_size}")
    print(f"  Elite solutions count: {elite_size}")
    print(f"  Mutation probability : {mutation_prob}")
    print(f"  Time limit (seconds) : {time_limit}")

    try:
        t, d, v, w, job_labels = read_job_data(Input_File_Name)
    except FileNotFoundError:
        print(f"Error: Input file '{Input_File_Name}.xlsx' was not found.")
        return

    n_jobs = len(t)
    run_start = time.perf_counter()

#הפעלת הסולבר על הסידור המקורי לקבלת ערך הבסיס
    original_sequence = list(range(n_jobs))
    baseline_score, _, _, _ = solve_lp_subproblem(original_sequence, t, d, v, w)

    best_record = {
        "sequence": original_sequence[:],
        "score": baseline_score,
        "generation": 0,
    }

    print("Stage 0: Baseline (original sequence) evaluated.")
    print(f"  Baseline objective value: {baseline_score}")

#אתחול והערכת דור ההורים הראשון
    population = create_initial_population(n_jobs, population_size)
    scores = evaluate_population(population, t, d, v, w)

    for individual, score in zip(population, scores):
        best_record = _update_best_solution(best_record, individual, score, 0)

    generation = 0
    last_reported_best = best_record["score"]
    last_improvement_generation = best_record["generation"]

    current_best_sequence, current_best_score = _population_best(population, scores)
    print("Stage 1: Initial population created and evaluated.")
    print(f"  Number of jobs       : {n_jobs}")
    print(f"  Best initial score   : {current_best_score}")

    #בניית דורות ברצף כל עוד לא הגענו למגבלת הזמן שנקבעה בקלט
    while time.perf_counter() - run_start < time_limit:
        generation += 1
        population, scores = build_next_generation(
            population,
            scores,
            t,
            d,
            v,
            w,
            population_size,
            elite_size,
            mutation_prob,
        )

        for individual, score in zip(population, scores):
            previous_best = best_record["score"]
            best_record = _update_best_solution(best_record, individual, score, generation)
            if (
                best_record["score"] is not None
                and previous_best is not None
                and best_record["score"] < previous_best
            ):
                last_improvement_generation = generation

        # הדפסת התקדמות לטרמינל כל 5 דורות 
        if generation % 5 == 0:
            print(f"Generation {generation}: Best objective value = {best_record['score']}")

            if (
                last_reported_best is not None
                and best_record["score"] is not None
                and best_record["score"] < last_reported_best
            ):
                print(
                    "  Improved since last report "
                    f"(new best found in generation {last_improvement_generation})."
                )

            last_reported_best = best_record["score"]

    run_time = time.perf_counter() - run_start

    print("GA run completed.")
    print(f"  Generations created  : {generation}")
    print(f"  Best objective value : {best_record['score']}")
    print(f"  Best found in gen.   : {best_record['generation']}")
    print(f"  Baseline objective   : {baseline_score}")
    print(f"  Run time (seconds)   : {run_time:.2f}")

# הרצה אחרונה של הסולבר עבור הרצף המנצח כדי לחלץ את זמני ההתחלה והסיום לפלט
    best_sequence = best_record["sequence"]
    _, completion_times, earliness, tardiness = solve_lp_subproblem(
        best_sequence, t, d, v, w
    )

    start_times = [completion_times[j] - t[j] for j in range(n_jobs)]
    penalties = [v[j] * earliness[j] + w[j] * tardiness[j] for j in range(n_jobs)]

    output_path = f"{Output_File_Na}.xlsx"
    try:
        # שליחה לפונקציית עיצוב האקסל הכוללת את בניית טבלת התוצאות והגאנט
        _write_results_workbook(
            output_path,
            job_labels,
            t,
            d,
            v,
            w,
            best_sequence,
            completion_times,
            start_times,
            earliness,
            tardiness,
            penalties,
            best_record["score"],
            baseline_score,
            run_time,
        )

        workbook_out = load_workbook(output_path)
        workbook_out.active["B14"] = generation
        workbook_out.save(output_path)
        workbook_out.close()

        print(f"Results saved to '{output_path}'.")
    except PermissionError:
        print(f"Error: Could not save '{output_path}'. Please close the file in Excel and run again.")
