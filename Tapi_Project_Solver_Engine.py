import time

import pulp
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter

# הגדרת צבעים לעיצוב תרשים הגאנט באקסל
JOB_FILL = PatternFill(start_color="556B2F", end_color="556B2F", fill_type="solid")
IDLE_FILL = PatternFill(start_color="D2B48C", end_color="D2B48C", fill_type="solid")
CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def solve_lp_subproblem(sequence, t, d, v, w):
#פתרון תכנית לינארית בעזרת ספריית PULP

    n = len(t)

    # הגדרת בעיית המינימום
    prob = pulp.LpProblem("JIT_Scheduling", pulp.LpMinimize)

    # הגדרת משתני החלטה רציפים
    C = [pulp.LpVariable(f"C_{i}", lowBound=0, cat="Continuous") for i in range(n)]
    E = [pulp.LpVariable(f"E_{i}", lowBound=0, cat="Continuous") for i in range(n)]
    T = [pulp.LpVariable(f"T_{i}", lowBound=0, cat="Continuous") for i in range(n)]

    # פונקציית המטרה מזעור סך הקנסות
    prob += pulp.lpSum(v[i] * E[i] + w[i] * T[i] for i in range(n))

    # אילוצי קדימויות וזמני עיבוד לפי הסדר שהתקבל
    if sequence:
        first_job = sequence[0]
        prob += C[first_job] >= t[first_job]

        for i in range(1, len(sequence)):
            curr_job = sequence[i]
            prev_job = sequence[i - 1]
            # משימה יכולה להסתיים רק אחרי שהמשימה הקודמת לה הסתיימה + זמן העיבוד שלה
            prob += C[curr_job] - t[curr_job] >= C[prev_job]

    # אילוצי הגדרת הקדמה ופיגור ביחס לזמן היעד 
    for j in range(n):
        prob += E[j] >= d[j] - C[j]
        prob += T[j] >= C[j] - d[j]

    # הגבלת זמן ריצה ל-5 שניות למניעת תקיעה במקרי קיצון
    prob.solve(pulp.PULP_CBC_CMD(timeLimit=5, msg=False))

    status = pulp.LpStatus[prob.status]
    if status != "Optimal":
        print(
            f"WARNING: LP solver finished with status '{status}' "
            f"(expected 'Optimal'). Check sequence or input data."
        )

    # חילוץ התוצאות מתוך משתני ההחלטה של הסולבר
    completion_times = [pulp.value(C[j]) for j in range(n)]
    earliness = [pulp.value(E[j]) for j in range(n)]
    tardiness = [pulp.value(T[j]) for j in range(n)]
    obj_value = pulp.value(prob.objective)

    return obj_value, completion_times, earliness, tardiness


def _fmt(value):
    # עיגול ערכים לספרה עשרונית אחת בשביל תצוגה בגאנט
    return round(float(value), 1)


def _read_data_row(ws, row, start_col=2):
    # פונקציית עזר לקריאת נתונים משורה באקסל
    values = []
    col = start_col
    while True:
        cell_val = ws.cell(row=row, column=col).value
        if cell_val is None or cell_val == "":
            break
        values.append(float(cell_val))
        col += 1
    return values


def _read_job_labels(ws, n_jobs, start_col=2):
    # קריאת אינדקסים של משימות מהשורה הראשונה בקובץ הקלט
    labels = []
    col = start_col
    while True:
        cell_val = ws.cell(row=1, column=col).value
        if cell_val is None or cell_val == "":
            break
        labels.append(int(cell_val))
        col += 1
    if len(labels) < n_jobs:
        labels = list(range(1, n_jobs + 1))
    return labels[:n_jobs]


def _read_input_from_workbook(ws):
    #  קריאת כלל הפרמטרים מהקלט (זמני עיבוד, יעדים וקנסות)
    t = _read_data_row(ws, 2)
    d = _read_data_row(ws, 3)
    v = _read_data_row(ws, 4)
    w = _read_data_row(ws, 5)
    n = len(t)
    job_labels = _read_job_labels(ws, n)
    return t, d, v, w, job_labels


def _build_schedule_columns(sequence, t, completion_times):
  #  בניית מבנה נתונים שמייצג את התזמון בפועל כולל איתור זמני בטלה
    columns = []

    if sequence:
        # בדיקה האם יש זמן בטלה לפני המשימה הראשונה
        first_job = sequence[0]
        first_start = completion_times[first_job] - t[first_job]
        if first_start > 1e-6:
            columns.append({"type": "idle", "start": 0.0, "end": first_start})

    # מעבר על שאר המשימות לאיתור זמני בטלה ביניהן
    for i, job_idx in enumerate(sequence):
        if i > 0:
            prev_job = sequence[i - 1]
            gap_start = completion_times[prev_job]
            gap_end = completion_times[job_idx] - t[job_idx]
            if gap_end > gap_start + 1e-6:
                columns.append(
                    {"type": "idle", "start": gap_start, "end": gap_end}
                )
        columns.append({"type": "job", "idx": job_idx})

    return columns


def _apply_column_fill(ws, col, fill, rows=(1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 17)):
    # צביעת עמודות באקסל להפרדה ויזואלית
    for row in rows:
        ws.cell(row=row, column=col).fill = fill


def _format_worksheet(ws):
    # התאמה של רוחב העמודות באקסל ויישור הטקסט למרכז
    for col_idx in range(1, ws.max_column + 1):
        max_len = 0
        for row_idx in range(1, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = CENTER_ALIGN
            if cell.value is not None:
                text = str(cell.value)
                line_len = max(len(line) for line in text.split("\n"))
                max_len = max(max_len, line_len)
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)


def _write_results_workbook(
    output_path,
    job_labels,
    t,
    d,
    v,
    w,
    sequence,
    completion_times,
    start_times,
    earliness,
    tardiness,
    penalties,
    obj_value,
    baseline_value,
    run_time,
):
    #  כתיבת התוצאות לקובץ אקסל
    schedule_columns = _build_schedule_columns(sequence, t, completion_times)

    wb_out = Workbook()
    ws = wb_out.active
    ws.title = "Results"

    # כתיבת כותרות לשורות
    row_labels = {
        1: "J", 2: "tj", 3: "dj", 4: "vj", 5: "wj", 
        6: "Sj", 7: "Cj", 8: "Ej", 9: "Tj", 10: "Pj",
    }
    for row, label in row_labels.items():
        ws.cell(row=row, column=1, value=label)

    col = 2
    for slot in schedule_columns:
        if slot["type"] == "idle":
            idle_start = slot["start"]
            idle_end = slot["end"]
            fs, fe = _fmt(idle_start), _fmt(idle_end)

            ws.cell(row=1, column=col, value=f"IDLE {fs}-{fe}")
            for row in (2, 3, 4, 5, 8, 9, 10):
                ws.cell(row=row, column=col, value="-")
            ws.cell(row=6, column=col, value=idle_start)
            ws.cell(row=7, column=col, value=idle_end)
            ws.cell(row=17, column=col, value=f"IDLE\n({fs} - {fe})")
            _apply_column_fill(ws, col, IDLE_FILL)
        else:
            job_idx = slot["idx"]
            lbl = job_labels[job_idx]
            fs = _fmt(start_times[job_idx])
            fc = _fmt(completion_times[job_idx])

            ws.cell(row=1, column=col, value=f"j{lbl}")
            ws.cell(row=2, column=col, value=t[job_idx])
            ws.cell(row=3, column=col, value=d[job_idx])
            ws.cell(row=4, column=col, value=v[job_idx])
            ws.cell(row=5, column=col, value=w[job_idx])
            ws.cell(row=6, column=col, value=start_times[job_idx])
            ws.cell(row=7, column=col, value=completion_times[job_idx])
            ws.cell(row=8, column=col, value=earliness[job_idx])
            ws.cell(row=9, column=col, value=tardiness[job_idx])
            ws.cell(row=10, column=col, value=penalties[job_idx])
            ws.cell(row=17, column=col, value=f"j{lbl}\n({fs} - {fc})")
            _apply_column_fill(ws, col, JOB_FILL)

        col += 1

    # הצגת השוואה וזמני ריצה בתחתית האקסל
    ws["A12"] = "Total Penalty"
    ws["B12"] = obj_value
    ws["A13"] = "Original Sol"
    ws["B13"] = baseline_value
    ws["A14"] = "Total Gen Created"
    ws["B14"] = "-"  # יתעדכן במידה והופעל האלגוריתם הגנטי
    ws["A15"] = "Run time"
    ws["B15"] = run_time
    ws["A17"] = "Gantt:"

    _format_worksheet(ws)
    wb_out.save(output_path)


def Run_Solver(Input_File_Name, Output_File_Name):
#    הפונקציה המרכזית להרצת חלק 1 בפרויקט
#    קוראת את הנתונים, מפעילה את הניסוח הלינארי על הסדר המקורי, ומייצאת לאקסל
    run_start = time.perf_counter()

    wb_in = load_workbook(f"{Input_File_Name}.xlsx", data_only=True)
    ws_in = wb_in.active
    t, d, v, w, job_labels = _read_input_from_workbook(ws_in)
    n = len(t)

    # יצירת סדר כרונולוגי לפי הקלט המקורי
    original_sequence = list(range(n))
    obj_value, completion_times, earliness, tardiness = solve_lp_subproblem(
        original_sequence, t, d, v, w
    )

    start_times = [completion_times[j] - t[j] for j in range(n)]
    penalties = [v[j] * earliness[j] + w[j] * tardiness[j] for j in range(n)]
    run_time = time.perf_counter() - run_start

    _write_results_workbook(
        f"{Output_File_Name}.xlsx",
        job_labels,
        t,
        d,
        v,
        w,
        original_sequence,
        completion_times,
        start_times,
        earliness,
        tardiness,
        penalties,
        obj_value,
        baseline_value=obj_value,
        run_time=run_time,
    )


if __name__ == "__main__":
    # בלוק בדיקה עצמית במידה ומריצים את הקובץ ישירות
    INPUT_NAME = "Input_Data_File_2_Example"
    OUTPUT_NAME = "My_Output_Results"

    print("מפעיל את הסולבר הלינארי ומייצר פלט...")
    try:
        Run_Solver(INPUT_NAME, OUTPUT_NAME)
        print(f"הצלחה! נוצר קובץ חדש בתיקייה בשם: {OUTPUT_NAME}.xlsx")
    except FileNotFoundError:
        print(f"שגיאה: לא מצאתי את קובץ הקלט '{INPUT_NAME}.xlsx'.")
        print("ודא שהקובץ של המרצה יושב באותה תיקייה של קובץ הפייתון.")
    except PermissionError:
        print("שגיאה: קובץ הפלט פתוח כרגע באקסל! סגור אותו והרץ שוב.")